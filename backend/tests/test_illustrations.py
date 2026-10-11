import hashlib
import io
import json
from unittest.mock import Mock
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from app.agent.state import create_initial_state
from app.agent.nodes import planner, designer, parser
from app.api.websocket import WorkflowRunner
from app.contracts.project_input import validate_spec
from app.models.page import PageState
from app.tools.photo_library import ROOT,photos,resolve,assign_images,crop_photo
from app.tools.slide_scene import build_scene,illustrated_layout,get_theme,render_scene_png
from app.tools.slide_scene import balanced_cover_title
from app.tools.pptx_generator import PPTXGenerator


def test_offline_library_has_at_least_40_unique_licensed_photos():
    library=photos(); assert len(library)>=40
    hashes=[]
    for pid,row in library.items():
        data=(ROOT/row['filename']).read_bytes(); digest=hashlib.sha256(data).hexdigest()
        assert digest==row['sha256']; hashes.append(digest)
        assert row['source_url'].startswith('https://commons.wikimedia.org/wiki/File:')
        assert row['author'] and row['license'] and row['license_url']
        with Image.open(io.BytesIO(data)) as image: assert image.width>=600 and image.height>=400
    assert len(hashes)==len(set(hashes))
    assert resolve('../../secrets') is None and resolve({'id':'forest-01'}) is None
    assert crop_photo('unknown',120,120) is None


def test_planner_selects_images_with_one_existing_llm_call(monkeypatch):
    ids=list(photos())[:2]
    outline=[dict(page_id=f'page-0{i+1}',page_type='cover' if i==0 else 'content',page_title='自然与生态',key_points=['保护森林'],suggest_layout='auto',image_id=ids[min(i,1)]) for i in range(4)]
    client=Mock(); client.generate.return_value=json.dumps(outline,ensure_ascii=False)
    monkeypatch.setattr(planner,'get_llm_client',lambda _:client)
    state=planner.planner_node(create_initial_state('photo',{'page_count':4,'auto_images':True},'森林保护'))
    assert not state['errors']; assert client.generate.call_count==1
    assert 'image_id' in client.generate.call_args.kwargs['system']
    assert sum(bool(p.image_id) for p in state['pages'].values())==2
    disabled=planner.generate_outline({'page_count':4,'auto_images':False})
    assert not any(item['image_id'] for item in disabled)
    invalid=assign_images([dict(page_type='content',image_id='../../secret')],True)
    assert invalid[0]['image_id'] is None
    assert assign_images([dict(page_type='content',image_id={'path':'secret'})],True)[0]['image_id'] is None
    assert validate_spec({'auto_images':False})=={'auto_images':False}


def test_picture_geometry_preview_export_and_attribution(tmp_path):
    pid=next(iter(photos())); row=photos()[pid]
    for kind in ('cover','content'):
        page=dict(page_id='page-01',page_type=kind,title='森林与未来',subtitle='让自然成为共同的责任',key_points=['守护生态','共同参与'],bullet_points=['保护多样性，珍惜每一片森林。','从身边的行动开始。'],image_id=pid,style_config={'layouts':[illustrated_layout(kind)]})
        scene=build_scene(page,get_theme('自然清新'))
        image=next(e for e in scene['elements'] if e['kind']=='image')
        for text in (e for e in scene['elements'] if e['kind']=='text'):
            assert min(image['x']+image['w'],text['x']+text['w'])<=max(image['x'],text['x']) or min(image['y']+image['h'],text['y']+text['h'])<=max(image['y'],text['y'])
        assert not any(e.get('overflow') for e in scene['elements'])
        png=tmp_path/(kind+'.png'); render_scene_png(scene,png)
        with Image.open(png) as preview:
            expected=crop_photo(pid,round(image['w']*120),round(image['h']*120))
            assert preview.getpixel((round(image['x']*120)+10,round(image['y']*120)+10))==expected.getpixel((10,10))
        ppt=tmp_path/(kind+'.pptx'); PPTXGenerator().generate([page],ppt,theme=get_theme('自然清新'))
        slide=Presentation(ppt).slides[0]
        assert sum(s.shape_type==MSO_SHAPE_TYPE.PICTURE for s in slide.shapes)==1
        assert row['source_url'] in slide.notes_slide.notes_text_frame.text and row['license'] in slide.notes_slide.notes_text_frame.text


def test_dense_or_disabled_content_keeps_text_and_no_photo():
    pid=next(iter(photos()))
    for enabled in (False,True):
        state=create_initial_state('dense',{'style':'商务简约','auto_images':enabled})
        state['current_page_id']='page-01'; state['outline']=[{'page_id':'page-01','suggest_layout':'auto'}]
        page=PageState(page_id='page-01',title='密集内容',image_id=pid,key_points=['说明'],bullet_points=['全部需要保留的事实与解释。'*150])
        state['pages']={'page-01':page}; designer.designer_node(state)
        assert page.image_id is None and page.bullet_points[0].endswith('全部需要保留的事实与解释。')
        assert not any(p['type']=='image' for p in page.layout_json['placeholders'])
    stored=PageState(page_id='p',image_id=pid).to_dict()
    assert WorkflowRunner._hydrate_page(stored).image_id==pid
    assert WorkflowRunner._normalize_page_for_export(stored)['image_id']==pid
    page={'page_type':'content','title':'照片已移除','image_id':'bad','style_config':{'layouts':[illustrated_layout('content')]}}
    assert not any(e['kind']=='image' for e in build_scene(page)['elements'])


def test_photo_cover_title_breaks_at_phrases_without_orphan_character():
    title='现代团队协作与办公空间'
    balanced=balanced_cover_title(title,6,42)
    assert balanced.replace('\n','')==title
    assert all(len(line)>=3 for line in balanced.split('\n'))
    explicit='现代团队\n办公协作'
    assert balanced_cover_title(explicit,6,42)==explicit
