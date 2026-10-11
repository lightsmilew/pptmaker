"""Exercise selectable designs through real geometry, exports and API boundaries."""
import pytest
from pptx import Presentation
from app.tools.design_catalog import STYLE_NAMES, catalog
from app.tools.slide_scene import build_scene, get_theme, template_layout
from app.tools.slide_quality import review_scene, contrast
from app.tools.pptx_generator import PPTXGenerator
from app.agent.nodes.designer import designer_node
from app.agent.nodes.writer import _generate_page_content
from app.agent.state import create_initial_state
from app.models.page import PageState, PageType
from fastapi.testclient import TestClient
from app.main import app


def page(index, count=3):
    return PageState(page_id=f"page-{index+1:02d}",page_type=PageType.CONTENT,
                     title="先理清结构，再展开观点",subtitle="让内容找到合适的位置",
                     key_points=[f"重点 {i+1}" for i in range(count)],
                     bullet_points=["保留材料中的事实，围绕听众需要展开解释。" for _ in range(count)])


@pytest.mark.parametrize("style", sorted(STYLE_NAMES))
def test_every_style_and_direction_preserves_content_and_readability(style):
    assert contrast(get_theme(style)["text"], get_theme(style)["background"]) >= 4.5
    for mode in ("auto","editorial","structured","spotlight"):
        state=create_initial_state("design-test",dict(style=style,layout_mode=mode))
        for i in range(5):
            p=page(i,i+1); state["pages"][p.page_id]=p
        for pid,p in state["pages"].items():
            state["current_page_id"]=pid; designer_node(state)
            assert not state["errors"]
            result=review_scene(p.to_dict(),style)
            assert result["status"] == "approved", (style,mode,pid,result["issues"])
            text="".join(e.get("text", "").replace("\n", "") for e in build_scene(p.to_dict(),get_theme(style))["elements"])
            assert all(detail in text for detail in p.bullet_points)


def test_generic_neighbouring_pages_have_distinct_layouts_without_llm(mock_llm_with_responses):
    client=mock_llm_with_responses([])
    state=create_initial_state("rhythm",dict(style="商务简约",layout_mode="auto"))
    for i in range(8):
        p=page(i); state["pages"][p.page_id]=p
    for pid in state["pages"]:
        state["current_page_id"]=pid; designer_node(state)
    layouts=[p.layout_type for p in state["pages"].values()]
    assert len(set(layouts)) >= 4
    assert all(a != b for a,b in zip(layouts,layouts[1:]))
    assert client.call_count == 0


def test_parallel_single_page_workers_keep_the_deck_order():
    outline=[dict(page_id=f"page-{i+1:02d}",suggest_layout="auto") for i in range(4)]
    layouts=[]
    for i,item in enumerate(outline):
        state=create_initial_state("parallel-rhythm",dict(style="商务简约",layout_mode="auto"))
        p=page(i); state["pages"][p.page_id]=p; state["current_page_id"]=p.page_id; state["outline"]=outline
        designer_node(state); layouts.append(p.layout_type)
    assert len(set(layouts)) == 4


def test_semantics_override_decoration_and_cover_families_differ():
    assert template_layout("content",3,True,"左右双栏",layout_mode="spotlight",page_index=1)["layout_type"] == "two_column"
    assert template_layout("content",3,True,"上下分区",layout_mode="structured")["layout_type"] == "timeline"
    assert template_layout("chart",3,True,layout_mode="editorial")["layout_type"] == "chart_content"
    covers=[template_layout("cover",family=get_theme(s)["family"])["cover_variant"] for s in ("商务简约","瑞士网格","科技风")]
    assert len(set(covers)) == 3


def test_varied_export_is_editable_and_uses_serif_title(tmp_path):
    state=create_initial_state("export",dict(style="杂志编辑",layout_mode="auto"))
    for i in range(4):
        p=page(i); state["pages"][p.page_id]=p
    for pid in state["pages"]:
        state["current_page_id"]=pid; designer_node(state)
    target=tmp_path/"varied.pptx"
    PPTXGenerator().generate([p.to_dict() for p in state["pages"].values()],str(target),theme=get_theme("杂志编辑"))
    prs=Presentation(target)
    assert len(prs.slides)==4
    assert len({tuple((shape.left,shape.top) for shape in s.shapes if shape.has_text_frame) for s in prs.slides}) >= 3
    assert any(p.font.name == "SimSun" for s in prs.slides for shape in s.shapes if shape.has_text_frame for p in shape.text_frame.paragraphs)
    assert not any(shape.shape_type == 13 for s in prs.slides for shape in s.shapes)


def test_density_changes_request_and_output_budget(mock_llm_with_responses):
    client=mock_llm_with_responses([{"items":[["重点","解释"]]}])
    _generate_page_content("page-02","content","主题","","项目",["重点"],content_density="airy")
    assert client.calls[0]["max_tokens"] == 700
    assert "每页2至3组" in client.calls[0]["prompt"]


def test_extra_model_groups_are_merged_without_losing_facts(mock_llm_with_responses):
    items=[[f"重点{i}",f"第{i}项原始解释"] for i in range(6)]
    mock_llm_with_responses([{"items":items}])
    result=_generate_page_content("page-02","content","标题","","主题",[],content_density="airy")
    assert len(result["key_points"]) == len(result["bullet_points"]) == 3
    for heading,detail in items:
        assert heading in "".join(result["key_points"])
        assert detail in "".join(result["bullet_points"])


def test_catalog_and_settings_contract(authenticated_client,tmp_path,monkeypatch):
    monkeypatch.setenv("FOLIO_DB",str(tmp_path/"catalog.sqlite3"))
    assert TestClient(app).get("/api/v1/design/catalog").status_code == 401
    client=authenticated_client()
    response=client.get("/api/v1/design/catalog")
    assert response.status_code == 200 and response.json() == catalog()
    assert len(response.json()["styles"]) == 16
    spec=dict(style="复古印刷",layout_mode="editorial",content_density="airy",page_count=6)
    created=client.post("/api/v1/projects",json={"spec":spec,"user_input":"文化分享"})
    assert created.status_code == 200 and created.json()["spec"] == {"auto_images": True, **spec}
    profile=client.patch("/api/v1/account/profile",json={"preferences":spec | {"page_count":6}})
    assert profile.status_code == 200
    assert profile.json()["profile"]["preferences"]["layout_mode"] == "editorial"
    for key,value in (("style","unknown"),("layout_mode","unknown"),("content_density","unknown")):
        assert client.post("/api/v1/projects",json={"spec":{key:value}}).status_code == 422
        assert client.patch("/api/v1/account/profile",json={"preferences":{key:value}}).status_code == 422
