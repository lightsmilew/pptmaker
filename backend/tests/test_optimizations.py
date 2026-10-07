import json
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation
from app.llm.config import LLMClient, LLMConfig, OpenAIConfig, _response_cache
from app.tools.slide_scene import build_scene, get_theme, render_scene_png
from app.tools.slide_quality import review_scene, contrast
from app.tools.pptx_generator import PPTXGenerator
from app.agent.nodes.writer import generate_page_content, writer_node
from app.agent.state import create_initial_state
from app.models.page import PageState, PageType


def sample_page():
    return dict(page_id="page-02", page_type="content", title="先整理结构，再讲清重点",
                subtitle="每组标题和解释都应保留在最终页面中",
                key_points=["明确主题", "建立顺序", "保留证据"],
                bullet_points=["说明演示的对象、听众和目的。", "按听众理解信息的顺序组织页面。", "保留原文数据，不补充没有依据的结论。"])


def test_shared_scene_preserves_details_and_editability(tmp_path):
    page = sample_page()
    scene = build_scene(page)
    text = "".join(e.get("text", "").replace("\n", "") for e in scene["elements"])
    assert all(detail in text for detail in page["bullet_points"])
    output = tmp_path / "sample.pptx"
    PPTXGenerator().generate([page],str(output))
    slide = Presentation(output).slides[0]
    exported = "".join(shape.text.replace("\n", "") for shape in slide.shapes if shape.has_text_frame)
    assert all(detail in exported for detail in page["bullet_points"])
    assert not any(shape.shape_type == 13 for shape in slide.shapes)
    preview = tmp_path / "preview.png"
    render_scene_png(scene, preview)
    assert Image.open(preview).size == (1600,900)


@pytest.mark.parametrize("style", ["商务简约", "学术简约", "科技风", "创意活泼", "教育培训", "金融专业"])
def test_templates_have_readable_contrast_and_geometry(style):
    theme = get_theme(style)
    assert contrast(theme["text"], theme["background"]) >= 4.5
    assert review_scene(sample_page(),style)["status"] == "approved"


def test_local_checks_detect_overflow_without_calling_llm():
    page = sample_page()
    page["bullet_points"] = ["非常长的详细说明"*100]*3
    result = review_scene(page)
    assert result["status"] == "needs_revision"
    assert result["review_method"] == "local_geometry"
    assert any(i["category"] == "overflow" for i in result["issues"])


def test_chart_and_table_are_native_objects(tmp_path):
    pages = [dict(page_id="page-01",page_type="chart",title="真实样本数据",
                  chart_data={"title":"样本", "data":{"labels":["甲","乙"],"values":[10,20]}}),
             dict(page_id="page-02",page_type="table",title="来源中的比较",
                  table_data={"headers":["对象","说明"],"rows":[["甲","文本"],["乙","其他文本"]]})]
    path=tmp_path/"evidence.pptx"
    PPTXGenerator().generate(pages,str(path))
    prs=Presentation(path)
    assert any(shape.has_chart for shape in prs.slides[0].shapes)
    assert any(shape.has_table for shape in prs.slides[1].shapes)


def test_compact_writer_contract_and_output_budget(mock_llm_with_responses):
    client = mock_llm_with_responses([{"items":[["具体主题","有依据的解释"]]}])
    result = generate_page_content("page-03","content","原有标题","原有副标题","主题",["要点"])
    assert result["title"] == "原有标题"
    assert result["key_points"] == ["具体主题"]
    assert result["bullet_points"] == ["有依据的解释"]
    assert client.calls[0]["max_tokens"] == 1000


def test_navigation_pages_skip_llm(mock_llm_with_responses):
    client = mock_llm_with_responses([])
    state = create_initial_state("cost-check",{"topic":"主题"})
    state["outline"]=[{"page_id":"page-01","page_type":"cover","page_title":"主题"}]
    state["current_page_id"]="page-01"
    state["pages"]["page-01"]=PageState(page_id="page-01",page_type=PageType.COVER)
    writer_node(state)
    assert client.call_count == 0


def test_response_cache_isolated_by_credentials_and_effective_settings(monkeypatch):
    _response_cache.clear()
    calls=[]
    monkeypatch.setattr(LLMClient,"_chat_openai",lambda self,messages,**kw: calls.append(kw) or "result")
    def client(key="a",temperature=.7):
        return LLMClient(LLMConfig(openai=OpenAIConfig(api_key=key,temperature=temperature)))
    messages=[{"role":"user","content":"same prompt"}]
    first=client()
    assert first.chat(messages,cache=True)==first.chat(messages,cache=True)
    assert len(calls)==1
    client("b").chat(messages,cache=True)
    client(temperature=.2).chat(messages,cache=True)
    assert len(calls)==3


def test_client_initialization_reuses_connection_under_concurrency(monkeypatch):
    import openai
    created=[]
    monkeypatch.setattr(openai,"OpenAI",lambda **kw: created.append(kw) or object())
    client=LLMClient(LLMConfig(openai=OpenAIConfig(api_key="fake")))
    with ThreadPoolExecutor(max_workers=4) as pool:
        clients=list(pool.map(lambda _:client._get_openai_client(),range(8)))
    assert len(created)==1
    assert all(item is clients[0] for item in clients)


def test_image_input_and_actual_usage_are_sent(monkeypatch):
    from app.llm.telemetry import usage_snapshot, project_context
    calls=[]
    usage=SimpleNamespace(prompt_tokens=1100,completion_tokens=100,
                          prompt_tokens_details=SimpleNamespace(cached_tokens=1024))
    response=SimpleNamespace(usage=usage,choices=[SimpleNamespace(finish_reason="stop",message=SimpleNamespace(content="{}"))])
    sdk=SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=lambda **kw:calls.append(kw) or response)))
    client=LLMClient(LLMConfig(openai=OpenAIConfig(api_key="fake")),role="critic")
    client._client=sdk
    token=project_context.set("image-test")
    try:
        client.generate_with_image("inspect", "png-data", "system",max_tokens=650)
    finally:
        project_context.reset(token)
    assert calls[0]["messages"][-1]["content"][1]["type"]=="image_url"
    metrics=usage_snapshot("image-test")
    assert metrics["output_tokens"]==100
    assert metrics["cached_input_tokens"]==1024
    assert metrics["provider_cache_hit_rate"]==pytest.approx(1024/1100)


def test_truncated_response_is_not_cached(monkeypatch):
    usage=SimpleNamespace(prompt_tokens=20,completion_tokens=10)
    response=SimpleNamespace(usage=usage,choices=[SimpleNamespace(finish_reason="length",message=SimpleNamespace(content="partial"))])
    client=LLMClient(LLMConfig(openai=OpenAIConfig(api_key="fake-truncated")))
    client._client=SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=lambda **kw:response)))
    _response_cache.clear()
    with pytest.raises(ValueError,match="truncated"):
        client.generate("test",cache=True)
    assert not _response_cache


def test_saved_page_round_trip_preserves_evidence_layout_and_review(tmp_path):
    from app.api.websocket import WorkflowRunner
    from app.models.page import PreviewInfo
    from app.agent.nodes.designer import generate_layout
    page = PageState(page_id="page-01",page_type=PageType.TABLE,title="对比",
                     key_points=["来源说明"],bullet_points=["具体解释必须保留"],
                     table_data={"headers":["对象","数据"],"rows":[["甲",10],["乙",20]]},
                     review_status="approved",review_score=100,review_method="local_geometry")
    config = generate_layout(page.page_id,"table",page.title,"",page.key_points,page.bullet_points)
    page.layout_json, page.master_json = config["layouts"][0],config["master"]
    page.preview = PreviewInfo(image_path=str(tmp_path / "preview.png"))
    project = {"pages":[{"page_id":page.page_id,"page_type":"table","status":"completed"}]}
    for save in (WorkflowRunner._save_page_content,WorkflowRunner._save_page_layout,
                 WorkflowRunner._save_page_preview,WorkflowRunner._save_page_review):
        save(project,0,page)
    restored = WorkflowRunner._hydrate_page(project["pages"][0])
    assert restored.table_data == page.table_data
    assert restored.layout_json == page.layout_json
    assert restored.preview == page.preview
    assert restored.review_method == "local_geometry"
    output=tmp_path/"saved-table.pptx"
    PPTXGenerator().generate([WorkflowRunner._normalize_page_for_export(project["pages"][0])],str(output))
    slide=Presentation(output).slides[0]
    assert any(shape.has_table for shape in slide.shapes)
    assert "具体解释必须保留" in "".join(s.text.replace("\n", "") for s in slide.shapes if s.has_text_frame)


def test_layout_repair_changes_actual_geometry_without_llm(mock_llm_with_responses):
    from app.agent.nodes.revision_planner import revision_planner_node
    from app.tools.slide_scene import template_layout
    client=mock_llm_with_responses([])
    page=PageState(**sample_page(),review_status="needs_revision",
                   review_issues=[{"owner":"layout","description":"元素超出画布"}])
    page.layout_json=template_layout("content",3,True)
    page.layout_json["placeholders"][0]["x"]=-1
    state=create_initial_state("repair-test",{"topic":"测试"})
    state["pages"][page.page_id]=page
    state["current_page_id"]=page.page_id
    revision_planner_node(state)
    assert not state["errors"]
    assert page.layout_type == "two_column"
    assert review_scene(page.to_dict())["status"] == "approved"
    assert client.call_count == 0


def test_deepseek_disables_reasoning_and_counts_native_cache(monkeypatch):
    from app.llm.telemetry import usage_snapshot, project_context
    monkeypatch.setenv("DEEPSEEK_REASONING_EFFORT","none")
    calls=[]
    usage=SimpleNamespace(prompt_tokens=300,completion_tokens=90,prompt_cache_hit_tokens=128,
                          completion_tokens_details=SimpleNamespace(reasoning_tokens=0))
    response=SimpleNamespace(usage=usage,choices=[SimpleNamespace(finish_reason="stop",message=SimpleNamespace(content='{"items":[]}'))])
    client=LLMClient(LLMConfig(openai=OpenAIConfig(api_key="fake",base_url="https://api.deepseek.com",model_name="deepseek-flash")),role="writer")
    client._client=SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=lambda **kw:calls.append(kw) or response)))
    token=project_context.set("deepseek-test")
    try:
        client.generate("return JSON",json_object=True,max_tokens=1000)
    finally:
        project_context.reset(token)
    assert calls[0]["reasoning_effort"] == "none"
    assert calls[0]["response_format"] == {"type":"json_object"}
    assert usage_snapshot("deepseek-test")["cached_input_tokens"] == 128


def test_missing_chart_uses_full_text_layout():
    page=sample_page()
    page["page_type"]="chart"
    scene=build_scene(page)
    assert not scene["evidence_rendered"]
    assert all(detail in "".join(e.get("text", "").replace("\n", "") for e in scene["elements"])
               for detail in page["bullet_points"])
    assert review_scene(page)["status"] == "approved"
