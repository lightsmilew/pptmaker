import io
import subprocess
from unittest.mock import patch

from PIL import Image
from pptx import Presentation

from .test_support_and_admin_presentations import support_clients, image_bytes
from .test_admin import clients
from app.tools.reference_images import extract_image_text


def test_reference_image_extraction_and_empty_image_brief(support_clients,monkeypatch):
    from app.tools import reference_images
    _,user,_,_=support_clients
    monkeypatch.setattr(reference_images,'extract_image_text',lambda data:('湿地保护\n生物多样性',''))
    result=user.post('/api/v1/upload/parse',files={'file':('screenshot.png',image_bytes(),'image/png')})
    assert result.status_code==200
    record=result.json(); assert '生物多样性' in record['content']
    created=user.post('/api/v1/projects',json={'uploaded_files':[{'upload_id':record['upload_id']}]})
    assert created.status_code==200 and '生物多样性' in created.json()['user_input']
    monkeypatch.setattr(reference_images,'extract_image_text',lambda data:('','未识别出文字'))
    blank=user.post('/api/v1/upload/parse',files={'file':('blank.png',image_bytes(),'image/png')}).json()
    assert blank['notice']=='未识别出文字'
    assert user.post('/api/v1/projects',json={'uploaded_files':[blank]}).status_code==422
    assert user.post('/api/v1/projects',json={'uploaded_files':[blank],'user_input':'湿地保护主题'}).status_code==200


def test_pptx_csv_references_and_invalid_images(support_clients):
    _,user,_,_=support_clients
    presentation=Presentation(); slide=presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text='复用既有演示'; slide.placeholders[1].text='提取这段参考内容'
    stream=io.BytesIO(); presentation.save(stream)
    ppt=user.post('/api/v1/upload/parse',files={'file':('reference.pptx',stream.getvalue(),'application/vnd.openxmlformats-officedocument.presentationml.presentation')})
    assert ppt.status_code==200 and '提取这段参考内容' in ppt.json()['content']
    csv=user.post('/api/v1/upload/parse',files={'file':('data.csv','月份,收入\n十月,120'.encode(),'text/csv')})
    assert csv.status_code==200 and '120' in csv.json()['content']
    assert user.post('/api/v1/upload/parse',files={'file':('fake.png',b'fake','image/png')}).status_code==400


def test_local_ocr_validation_output_and_timeout(monkeypatch):
    monkeypatch.setenv('TESSERACT_CMD','isolated-test-tesseract')
    with patch('app.tools.reference_images.subprocess.run',return_value=subprocess.CompletedProcess([],0,'中文 English'.encode(),b'')) as run:
        text,notice=extract_image_text(image_bytes())
        assert text=='中文 English' and not notice
        assert run.call_args.kwargs['timeout']==30
        assert 'chi_sim+eng' in run.call_args.args[0]
    with patch('app.tools.reference_images.subprocess.run',side_effect=subprocess.TimeoutExpired('ocr',30)):
        text,notice=extract_image_text(image_bytes())
        assert not text and '未完成' in notice
