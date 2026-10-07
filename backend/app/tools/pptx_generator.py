"""Editable PowerPoint export from the shared, measured slide scene."""

from pathlib import Path
from .slide_scene import build_scene, normalize_page, get_theme, rgb, PALETTES


class PPTXGenerator:
    def __init__(self):
        try:
            import pptx
            self.pptx_available = True
        except ImportError:
            self.pptx_available = False

    def generate(self, pages, output_path, title="PPT Presentation", theme=None, template_path=None):
        if not self.pptx_available:
            raise RuntimeError("python-pptx is required to export a valid PPTX")
        if template_path:
            raise ValueError("External PPTX masters are not supported; select a built-in theme")
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        prs._element.sldSz.set("type", "custom")
        prs.core_properties.title = title
        for page in pages:
            data = normalize_page(page)
            scene = build_scene(data, theme or get_theme())
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            self._render_scene(slide, scene)
            if data.get("speaker_notes"):
                slide.notes_slide.notes_text_frame.text = data["speaker_notes"]
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_path)
        return output_path

    def _render_style_config(self, slide, page, config, theme=None):
        self._render_scene(slide, build_scene({**page, "style_config": config}, theme or get_theme()))

    def _render_scene(self, slide, scene):
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(*rgb(scene["background"]))
        for e in scene["elements"]:
            x,y,w,h = (Inches(e[k]) for k in ("x","y","w","h"))
            if e["kind"] == "rect":
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,x,y,w,h)
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor(*rgb(e["color"]))
                shape.line.fill.background()
            elif e["kind"] == "text":
                box = slide.shapes.add_textbox(x,y,w,h)
                frame = box.text_frame
                frame.word_wrap = False
                frame.margin_left = frame.margin_right = 0
                frame.margin_top = frame.margin_bottom = 0
                frame.vertical_anchor = MSO_ANCHOR.TOP
                for i,line in enumerate(e["text"].split("\n")):
                    p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                    p.text = line
                    p.font.name = "Microsoft YaHei"
                    p.font.size = Pt(e["size"])
                    p.font.bold = e["bold"]
                    p.font.color.rgb = RGBColor(*rgb(e["color"]))
                    p.space_before = p.space_after = Pt(0)
                    p.line_spacing = Pt(e["size"]*e["line_height"])
            elif e["kind"] == "chart":
                from pptx.chart.data import CategoryChartData
                from pptx.enum.chart import XL_CHART_TYPE
                payload = e["data"]["data"]
                data = CategoryChartData()
                data.categories = payload["labels"]
                data.add_series(e["data"].get("title", "数据"), payload["values"])
                chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,x,y,w,h,data).chart
                chart.has_legend = False
                chart.chart_style = 10
                series = chart.series[0]
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = RGBColor(*rgb(e["accent"]))
                series.format.line.fill.background()
                for axis in (chart.category_axis, chart.value_axis):
                    axis.tick_labels.font.name = "Microsoft YaHei"
                    axis.tick_labels.font.size = Pt(12)
                    axis.tick_labels.font.color.rgb = RGBColor(*rgb(e["color"]))
                chart.value_axis.has_major_gridlines = False
            elif e["kind"] == "table":
                rows = [e["data"]["headers"]] + e["data"]["rows"]
                table = slide.shapes.add_table(len(rows),len(rows[0]),x,y,w,h).table
                for ri,row in enumerate(rows):
                    for ci,value in enumerate(row):
                        cell = table.cell(ri,ci)
                        cell.text = str(value)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor(*rgb(e["accent"] if ri == 0 else e["background"]))
                        cell.margin_top = cell.margin_bottom = Inches(.1)
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.name = "Microsoft YaHei"
                            paragraph.font.size = Pt(16)
                            paragraph.font.bold = ri == 0
                            paragraph.font.color.rgb = RGBColor(*rgb(e["header_color"] if ri == 0 else e["color"]))


class TemplateLibrary:
    TEMPLATES = {name: get_theme(name) for name in PALETTES}

    def get_template(self, name):
        return get_theme(name)

    def list_templates(self):
        return list(self.TEMPLATES)
